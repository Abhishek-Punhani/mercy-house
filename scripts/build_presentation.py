"""Build the editable dark technical submission deck and its PDF.

Authoring only: python-pptx 1.0.2 and LibreOffice, never runtime dependencies.
Unchanged input manifests reuse the already-rendered artifacts.
"""
from pathlib import Path
import hashlib
import importlib.util
import json
import subprocess
import tempfile
import shutil
import re

ROOT=Path(__file__).resolve().parents[1]
DOC=ROOT/'docs'
INPUTS=['scripts/build_presentation.py','docs/ui-report.png','docs/ui-interview.png',
        'docs/SOURCE_TRAINING_RESULT.json','docs/ELITE_TRANSFER_RESULT.json',
        'docs/CASE_SOLVABILITY_AUDIT.json','docs/VALIDATION_RESULT.json',
        'docs/FAMILY_VALIDATION_RESULT.json']

def read(name,default=None):
    p=DOC/name
    return json.loads(p.read_text()) if p.exists() else (default or {})


def build():
    fingerprint={s:hashlib.sha256((ROOT/s).read_bytes()).hexdigest() if (ROOT/s).exists() else None for s in INPUTS}
    manifest=DOC/'PRESENTATION_MANIFEST.json'
    if manifest.exists() and read(manifest.name).get('inputs_sha256')==fingerprint and all((DOC/f).exists() for f in ('deck.pptx','deck.pdf')):
        print(json.dumps({'presentation':'unchanged','slides':12}));return
    if importlib.util.find_spec('pptx') is None:
        raise SystemExit('Presentation inputs changed. Install authoring dependencies with pip install -r requirements-presentation.txt, then rerun scripts/build_presentation.py.')
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from PIL import Image
    R=Presentation();R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
    C={'bg':'090F1B','panel':'111D30','edge':'24354B','ink':'EDF4FF','muted':'A4B6CD','mint':'66E6CD','amber':'FFC177','violet':'A2ADFF','red':'FF8798'}
    texts=[]
    def color(k):return RGBColor.from_string(C.get(k,k))
    def box(s,x,y,w,h,fill='panel',line=None,radius=True):
        shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=color(fill)
        if line:shape.line.color.rgb=color(line)
        else:shape.line.fill.background()
        if radius:
            try:shape.adjustments[0]=.08
            except IndexError:pass
        return shape
    def text(s,x,y,w,h,words,size=18,fg='ink',bold=False,font='DejaVu Sans'):
        sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.015);tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(str(words).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fg);p.space_after=Pt(8)
        return sh
    def slide(title,kicker,subtitle=''):
        s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color('bg')
        for x in (8.3,9.5,10.7,11.9):box(s,x,0,.012,7.5,'panel',radius=False)
        box(s,.48,.38,.10,.16,'mint',radius=False);text(s,.67,.33,11,.24,'MERCY HOUSE  /  '+kicker.upper(),10,'mint',True)
        text(s,.6,.94,12.1,.9,title,31,'ink',True)
        if subtitle:text(s,.62,1.85,12,.65,subtitle,15,'muted')
        box(s,.6,7.03,12.1,.012,'edge',radius=False)
        text(s,.62,7.15,11.3,.22,'KASUKABE DEFENCE FORCE  •  RL ENVIRONMENT HACKATHON',8,'muted')
        text(s,12,7.12,.7,.25,f'{len(R.slides):02d} / 12',9,'mint')
        texts.append({'title':title,'kicker':kicker,'subtitle':subtitle});return s
    def card(s,x,y,w,h,label,body,accent='mint',body_size=17):
        box(s,x,y,w,h,line='edge');box(s,x+.18,y+.2,.05,.34,accent,radius=False)
        text(s,x+.36,y+.23,w-.55,.45,label,16,accent,True);text(s,x+.23,y+.92,w-.46,h-1.05,body,body_size)
    def image(s,name,x,y,w,h):
        p=DOC/name
        if not p.exists():return
        iw,ih=Image.open(p).size;scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
        s.shapes.add_picture(str(p),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
    def note(s,words):s.notes_slide.notes_text_frame.text=words
    def small(s,words):text(s,.65,6.54,12,.35,words,11,'muted')
    train=read('SOURCE_TRAINING_RESULT.json');elite=read('ELITE_TRANSFER_RESULT.json');audit=read('CASE_SOLVABILITY_AUDIT.json');tests=read('VALIDATION_RESULT.json')
    # 1
    s=slide('Truth needs a chain of evidence.','Evidence-grounded agent evaluation')
    text(s,.65,2.05,6.0,1.15,'An investigation world for language agents.',29,'mint',True)
    text(s,.65,3.36,5.7,1.15,'A snowbound town in 1900. Private lives, limited witnesses and evidence that can point at the wrong person.',20)
    image(s,'ui-interview.png',6.8,2.05,5.85,4.1)
    text(s,6.95,6.19,5.6,.25,'Shown: scripted inquest reference, not a live solve.',10,'muted')
    for x,n,l in [(.65,'8','PUBLIC TASKS'),(2.6,'4','CASE FAMILIES'),(4.55,'3','DESIGNED LEVELS')]:
        box(s,x,5.15,1.7,1.12,line='edge');text(s,x+.16,5.24,1.3,.48,n,28,'mint',True);text(s,x+.16,5.86,1.4,.3,l,8,'muted')
    small(s,'The competition deliverable is the environment. Judges supply their detective model.')
    # 2
    s=slide('The agent controls the investigation.','What the environment asks an agent to do','It starts with its own biography and the local scene. The answer and hidden event journal remain private.')
    for i,(a,b) in enumerate([('01  EXPLORE','Choose a place and inspect visible features.'),('02  INTERVIEW','Question people; separate experience from hearsay.'),('03  TEST','Compare acquired documents and physical traces.'),('04  REPORT','Cite sources for each conclusion; retain uncertainty.')]):card(s,.65+i*3.08,2.7,2.9,2.55,a,b,body_size=16)
    text(s,.72,5.65,11.9,.72,'Beginner: one witnessed assault. Intermediate: two casualties and an intended target. Advanced: fabrication, custody, historical records and divided loyalties.',17,'muted')
    small(s,'Eight task scopes across four causal families; difficulty labels describe design complexity, not measured success bands.')
    # 3
    s=slide('The incident executes before arrival.','World state and causal traces','Validated transitions produce the injury, objects and local recollections that the detective can investigate.')
    labels=[('ACTION TAPE','Recorded or authored\nseeded decisions'),('WORLD ENGINE','Movement, possession,\nattack and delivery'),('LOCAL EFFECTS','Physical traces +\nlimited perception'),('INVESTIGATION','Persistent tools,\nreceipts and dialogue')]
    for i,(a,b) in enumerate(labels):
        card(s,.65+3.08*i,2.75,2.9,2.05,a,b,body_size=15)
    for x,a,b in [(.65,'NO ASSAULT','No fatal lamp injury.'),(4.8,'INTACT BRAKE','No sabotage-caused coach deaths.'),(8.95,'ORDINARY CORDIAL','No adulteration-caused death.')]:card(s,x,5.03,3.75,1.23,a,b,'amber',12)
    small(s,'Mercy House replays recorded LLM actions. New cases execute authored seeded tapes. Continuous offstage lives are unfinished.')
    # 4
    s=slide('Each character owns its memories.','Character model architecture','Private planning and public speech use separate contexts; physical truth stays in the engine.')
    card(s,.65,2.75,3.7,2.6,'PRIVATE PLANNER','Own life, local memories and legal choices. The raw detective question is replaced by derived topics.',body_size=17)
    card(s,4.82,2.75,3.7,2.6,'DISCLOSURE GATE','Trust, relevant acquired evidence and memory sensitivity determine which recollections may be released.','violet',17)
    card(s,8.99,2.75,3.7,2.6,'PUBLIC SPEAKER','Only released recollections and public context become the spoken answer. Dialogue cannot create objects.','amber',17)
    text(s,.72,5.74,11.9,.6,'The viewer streams observed actions, actual dialogue and stated notes. It excludes private character memories and provider reasoning.',17,'muted')
    # 5
    s=slide('The submitted report stays visible.','Observable results and official scoring')
    image(s,'ui-report.png',.65,2.0,7.2,4.55)
    box(s,8.05,2.15,4.6,3.98,line='edge');text(s,8.3,2.4,4.1,.4,'PROGRAMMATIC VERIFIER',15,'mint',True)
    text(s,8.3,3.02,4.05,.77,'R = clamp(C − .20F − .10U, 0, 1)',21,'ink',True)
    text(s,8.3,4.03,4.02,1.87,'C: supported target coverage\nF: distinct false findings\nU: distinct unsupported findings\nA correct guess without proof earns zero.',16)
    small(s,'Screenshot: actual native detective success, Lantern Room seed211. Narrative is unscored. Optional author-review LLM cannot assign rewards.')
    # 6
    s=slide('A small stack with explicit boundaries.','Implementation and isolation','Python standard-library runtime • Vanilla JavaScript, CSS and original SVG • Docker Compose • HTTPS model APIs')
    for x,a,b,accent in [(.65,'PROTECTED WORLD','Python state engine, receipts, proof routes and deterministic verifier.','mint'),(3.73,'TRUSTED HARNESS','Model-independent parsing, step/time budgets, private HTTP tool requests and JSONL logs.','violet'),(6.81,'MODEL GATEWAY','Configured endpoints, credential isolation, deadlines and request limits.','amber'),(9.89,'AGENT WORKSPACE','Real shell plus public town CLI; no keys, truth files or external route.','mint')]:card(s,x,2.75,2.8,2.65,a,b,accent,15)
    text(s,.72,5.75,11.9,.64,'Presentation UI: local Python HTTP server + server-sent events. Policy training: CPU REINFORCE over a linear softmax selector; no GPU or LLM fine-tuning.',17,'muted')
    small(s,'Workspace has no Docker socket, private mounts or control token. Security assumes a trusted host and Docker daemon.')
    # 7
    s=slide('RL learns which proposed action to execute.','The trained policy','Optional research policy: a compact learned selector ranks actions using public features. The native UI/judging detective remains independent.')
    for i,(a,b) in enumerate([('OBSERVED STATE','Acquired records, comparisons, interviews, repeat attempts and steps left.'),('CANDIDATE ACTIONS','An authored public proposer during CPU training; an LLM proposer during transfer tests.'),('LEARNED SELECTOR','Scores features, produces action probabilities, and executes one choice.')]):card(s,.65+4.13*i,2.75,3.91,2.65,a,b,['mint','violet','amber'][i],17)
    text(s,.73,5.8,11.85,.57,'REINFORCE raises the probability of action sequences that earn a valid terminal exercise reward, relative to a running baseline.',18,'muted')
    small(s,'Earlier prototype:20,000 episodes /460 weights. Fresh source curriculum:6,000 episodes /132 weights. Language-model weights stay fixed.')
    # 8
    s=slide('Three mechanisms become practice exercises.','Sourced inspiration, original interactive training','Original renamed checkpoint scenarios reuse the tested world mechanics. They do not reproduce TV scripts or complete episodes.')
    cards=[('INTENDED TARGET','The Mentalist S1E8\nSeparate the person targeted from the other casualty.'),('APPARENT FOREKNOWLEDGE','The Mentalist S1E7\nTest authorship; prediction alone does not establish murder.'),('TRUSTED RELATIVE','His & Hers\nCheck a familiar parent against evidence rather than loyalty or rumour.')]
    for i,(a,b) in enumerate(cards):card(s,.65+4.13*i,2.75,3.91,2.6,a,b,['mint','violet','amber'][i],17)
    text(s,.72,5.78,11.9,.62,'Both culprit assignments occur. Training starts from zero; the elite Mercy House cases are excluded from this new curriculum.',17,'muted')
    small(s,'Sources credited in SOURCE_TRAINING_REPORT.md. Some materials are acquired before these checkpoints begin; this is skill training.')
    note(s,'Sources: The Mentalist S1E7 Seeing Red https://www.tvmaze.com/episodes/8884/the-mentalist-1x07-seeing-red ; S1E8 The Thin Red Line https://tv.apple.com/nz/episode/thin-red-line/umc.cmc.343dcf9kznkxsuz2rx6jcazs9 ; His & Hers official ending https://www.netflix.com/tudum/articles/his-hers-episode-6-recap-ending-explained . Broad mechanisms only; no episode transcript training or original-plot generalization claim.')
    # 9
    s=slide('Reward supported inference, not activity.','New CPU training and held-out evaluation')
    episodes=train.get('training_episodes',train.get('episodes','Pending'))
    evaluations=train.get('evaluations',{})
    text(s,.7,2.18,11.8,.62,f'{episodes} episodes  /  48 training checkpoints  /  48 disjoint held-out checkpoints  /  CPU only',20,'mint',True)
    card(s,.65,3.0,5.25,2.67,'TERMINAL EXERCISE REWARD','1: the scoped material finding has acquired proof and no false or unsupported conclusions.\n0: incomplete, guessed, wrong or unsupported.','mint',17)
    labels=[('Untrained','untrained'),('Trained sampling','trained_sampling'),('Trained greedy','trained_greedy')]
    for i,(label,key) in enumerate(labels):
        data=evaluations.get(key,{})
        n=data.get('episodes',0);rate=data.get('mean_reward')
        if rate is None and isinstance(data.get('success_rate'),(int,float)):rate=data['success_rate']
        y=3.05+i*1.0;text(s,6.35,y,4.7,.3,label,16,'muted')
        box(s,6.35,y+.42,4.8,.18,'edge',radius=False)
        if isinstance(rate,(int,float)):box(s,6.35,y+.42,max(.03,4.8*rate),.18,'mint' if i else 'amber',radius=False)
        text(s,11.35,y+.23,1.2,.38,f'{rate:.1%}' if isinstance(rate,(int,float)) else 'Pending',17,'ink',True)
    small(s,'96 trials per policy (48 checkpoints repeated twice). Prepared acquisitions + authored candidate rules. Same templates; not full-case success.')
    # 10
    s=slide('Test transfer on excluded elite cases.','Frozen checkpoint • matched proposal contract','Two predeclared task scopes of the excluded Mercy House incident. LLM-ranked baseline versus frozen source-trained selector;40 choices per arm.')
    rows=elite.get('rows',[])
    for i,task in enumerate(['last_witness','mercy_inquest']):
        y=2.8+i*1.48;box(s,.65,y,12.03,1.2,line='edge');text(s,.9,y+.2,4.5,.4,'Fatal assault' if i==0 else 'Seven-act full inquest',21,'ink',True)
        for j,arm in enumerate(['llm_ranked','rl_selected']):
            row=next((r for r in rows if r.get('task_id')==task and r.get('arm')==arm),{})
            result=row.get('result',{});score=result.get('score')
            text(s,5.9+j*3.25,y+.14,3.05,.27,'LLM-RANKED' if j==0 else 'SOURCE-TRAINED SELECTOR',10,'muted',True)
            text(s,5.9+j*3.25,y+.51,3.05,.48,str(score) if score is not None else 'Pending',27,'mint' if result.get('success') else 'amber',True)
    text(s,.78,5.95,11.85,.4,'The same grounded response bank controls NPC variation. After branches diverge, later LLM proposals may differ.',15,'muted')
    small(s,'Small development comparison, not statistical superiority or the standard native-tool judging harness. Full trajectories retained.')
    # 11
    s=slide('Audit the case, not just the score.','Consistency, source-only solvability and robustness')
    count=tests.get('town_tests','?');node=tests.get('node_state_tests',8)
    card(s,.65,2.75,3.85,2.85,'FACT AND PROVENANCE FIXES','Distinguish payer from carrier. Authenticate the register writer independently. Correct what witnesses actually perceived.','mint',17)
    card(s,4.76,2.75,3.85,2.85,'LEGAL PROOF ROUTES', 'All eight reference/no-op pairs: 1 / 0. Public-information-only scripted probes solve all three new families with both culprit roles.','violet',17)
    card(s,8.87,2.75,3.85,2.85,'EXECUTABLE CHECKS',f'{count} town + 9 legacy + {node} Node tests. Browser desktop/mobile reports, isolation probes, forged-source rejection and causal counterfactuals.','amber',17)
    small(s,'Scripted feasibility is not live cooperation. Canonical murder proof can depend on a reluctant witness or a supported admission.')
    # 12
    s=slide('A reviewable environment, with measured limits.','Demonstration and handoff')
    card(s,.65,2.7,5.82,2.83,'TRY THE INVESTIGATION','localhost:8787 → New investigation → Lantern Room, seed7 → configured key.\nThen inspect the submitted report and acquired citations.','mint',18)
    card(s,6.77,2.7,5.82,2.83,'WHAT IS ESTABLISHED','Four native beginner runs scored 1.0 with live character agents. Harder-case and RL results are reported separately. Continuous autonomous lives and scored ghost causes remain unfinished.','amber',17)
    text(s,.75,5.83,11.88,.5,'docker compose up --build --abort-on-container-exit --exit-code-from harness',16,'violet',font='DejaVu Sans Mono')
    small(s,'Source + baseline gates + raw runs + editable deck/PDF + security and audit evidence. Final character-model setup and submission are external.')
    for s in R.slides:
        if not s.notes_slide.notes_text_frame.text.strip():note(s,'All claims describe implemented local artifacts and recorded development measurements. See docs/PS_COMPLIANCE_AUDIT.md, SOURCE_TRAINING_REPORT.md, ELITE_TRANSFER_RESULT.json and CASE_SOLVABILITY_AUDIT.md. Do not present scripted references as LLM success or checkpoint rates as whole-case success.')
    R.save(str(DOC/'deck.pptx'))
    with tempfile.TemporaryDirectory(prefix='mercy-deck-render-') as folder:
        profile=Path(folder)/'profile';output=Path(folder)/'pdf';output.mkdir()
        with (ROOT/'logs/presentation-render.log').open('w') as log:
            subprocess.run(['libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(output),str(DOC/'deck.pptx')],check=True,timeout=90,stdout=log,stderr=subprocess.STDOUT)
        shutil.copyfile(output/'deck.pdf',DOC/'deck.pdf')
    info=subprocess.check_output(['pdfinfo',str(DOC/'deck.pdf')],text=True);pages=int(re.search(r'^Pages:\s+(\d+)',info,re.M)[1]);assert pages==12,pages
    manifest.write_text(json.dumps({'slides':12,'theme':'dark technical; editable native shapes + actual UI screenshots','inputs_sha256':fingerprint,'screenshots':{'ui-report.png':'actual native detective Lantern Room seed211 success','ui-interview.png':'explicitly labeled original inquest scripted reference'}},indent=2)+'\n')
    import html
    (DOC/'deck.html').write_text('<!doctype html><meta charset="utf-8"><title>Mercy House presentation</title><style>body{background:#090f1b;color:#edf4ff;font:18px sans-serif;max-width:900px;margin:60px auto}a{color:#66e6cd}li{margin:20px}</style><h1>Mercy House</h1><p>Editable dark technical presentation: <a href="deck.pptx">PowerPoint</a> · <a href="deck.pdf">PDF</a></p><ol>'+''.join('<li>'+html.escape(t['title'])+'</li>' for t in texts)+'</ol>')
    print(json.dumps({'slides':pages,'pptx':str(DOC/'deck.pptx'),'pdf':str(DOC/'deck.pdf')}))

if __name__=='__main__':build()
